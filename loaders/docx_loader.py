from pathlib import Path
from typing import List
from langchain_core.documents import Document

def try_unstructured_docx(path: Path):
    try:
        from unstructured.partition.docx import partition_docx
        els = partition_docx(filename=str(path))
        txt = "\n".join(e.text for e in els if hasattr(e, "text"))
        if txt.strip():
            return [Document(page_content=txt, metadata={"source": str(path)})]
    except:
        pass
    return []

def try_python_docx(path: Path):
    try:
        import docx
        d = docx.Document(str(path))
        txt = "\n".join([p.text for p in d.paragraphs if p.text])
        return [Document(page_content=txt, metadata={"source": str(path)})]
    except:
        return []

def load_docx(files: List[Path]) -> List[Document]:
    out = []
    for f in files:
        # Tries unstructured first; falls back to python-docx if first fails
        d = try_unstructured_docx(f) or try_python_docx(f)
        out.extend(d)
    return out